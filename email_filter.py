"""
Namecheap Email Filter & Organizer
-----------------------------------
Connects to your Namecheap email via IMAP, categorizes each email using
configurable rules (keywords, sender domains, subject patterns), and exports
organized data to an Excel (.xlsx) file. No API key required.

Usage:
    python email_filter.py

    Optional flags:
        --limit N       Process only the N most recent emails (default: 100)
        --folder FOLDER IMAP folder to read from (default: INBOX)
        --output FILE   Output Excel filename (default: email_report.xlsx)
        --since DAYS    Only fetch emails from the last N days (default: 30)
        --rules FILE    Path to rules JSON file (default: email_rules.json)
"""

import argparse
import email
import imaplib
import json
import os
import re
import ssl
import sys
from datetime import datetime, timedelta
from email.header import decode_header
from email.utils import parsedate_to_datetime

import openpyxl
from dotenv import load_dotenv
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

load_dotenv()

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

IMAP_HOST = os.getenv("EMAIL_IMAP_HOST", "mail.privateemail.com")
IMAP_PORT = int(os.getenv("EMAIL_IMAP_PORT", "993"))
EMAIL_ADDRESS = os.getenv("EMAIL_ADDRESS", "")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD", "")

# ---------------------------------------------------------------------------
# IMAP helpers
# ---------------------------------------------------------------------------


def connect_imap() -> imaplib.IMAP4_SSL:
    """Open an authenticated IMAP connection to the mail server."""
    if not EMAIL_ADDRESS or not EMAIL_PASSWORD:
        sys.exit(
            "ERROR: EMAIL_ADDRESS and EMAIL_PASSWORD must be set in your .env file."
        )
    context = ssl.create_default_context()
    conn = imaplib.IMAP4_SSL(IMAP_HOST, IMAP_PORT, ssl_context=context)
    conn.login(EMAIL_ADDRESS, EMAIL_PASSWORD)
    return conn


def decode_str(value) -> str:
    """Decode an encoded email header value to a plain string."""
    if value is None:
        return ""
    if isinstance(value, bytes):
        try:
            return value.decode("utf-8", errors="replace")
        except Exception:
            return str(value)
    parts = decode_header(value)
    decoded = []
    for part, charset in parts:
        if isinstance(part, bytes):
            decoded.append(part.decode(charset or "utf-8", errors="replace"))
        else:
            decoded.append(part)
    return "".join(decoded)


def get_body(msg: email.message.Message) -> str:
    """Extract plain-text body from an email message (first 2 000 chars)."""
    body = ""
    if msg.is_multipart():
        for part in msg.walk():
            ctype = part.get_content_type()
            disposition = str(part.get("Content-Disposition", ""))
            if ctype == "text/plain" and "attachment" not in disposition:
                payload = part.get_payload(decode=True)
                if payload:
                    charset = part.get_content_charset() or "utf-8"
                    body = payload.decode(charset, errors="replace")
                    break
    else:
        payload = msg.get_payload(decode=True)
        if payload:
            charset = msg.get_content_charset() or "utf-8"
            body = payload.decode(charset, errors="replace")
    return body[:2000].strip()


def fetch_emails(
    conn: imaplib.IMAP4_SSL,
    folder: str = "INBOX",
    limit: int = 100,
    since_days: int = 30,
) -> list:
    """Fetch the most recent emails from a folder and return as list of dicts."""
    conn.select(folder, readonly=True)

    since_date = (datetime.now() - timedelta(days=since_days)).strftime("%d-%b-%Y")
    status, data = conn.search(None, f'(SINCE "{since_date}")')
    if status != "OK":
        print(f"No messages found in {folder}.")
        return []

    message_ids = data[0].split()
    # Most recent first
    message_ids = message_ids[-limit:][::-1]
    print(f"Fetching {len(message_ids)} emails from {folder}...")

    emails = []
    for i, msg_id in enumerate(message_ids, 1):
        status, msg_data = conn.fetch(msg_id, "(RFC822)")
        if status != "OK":
            continue
        raw = msg_data[0][1]
        msg = email.message_from_bytes(raw)

        subject = decode_str(msg.get("Subject", "(no subject)"))
        sender = decode_str(msg.get("From", ""))
        date_str = msg.get("Date", "")
        try:
            date_obj = parsedate_to_datetime(date_str)
            date_formatted = date_obj.strftime("%Y-%m-%d %H:%M")
        except Exception:
            date_formatted = date_str

        body = get_body(msg)

        emails.append(
            {
                "id": msg_id.decode(),
                "subject": subject,
                "sender": sender,
                "date": date_formatted,
                "body": body,
            }
        )

        if i % 10 == 0:
            print(f"  Fetched {i}/{len(message_ids)}...")

    return emails


# ---------------------------------------------------------------------------
# Rule-based categorization
# ---------------------------------------------------------------------------


def load_rules(rules_path: str) -> dict:
    """Load categorization rules from a JSON file."""
    if not os.path.exists(rules_path):
        sys.exit(
            f"ERROR: Rules file not found: {rules_path}\n"
            "Copy email_rules.json.example to email_rules.json and customize it."
        )
    with open(rules_path, "r", encoding="utf-8") as f:
        return json.load(f)


def _extract_sender_domain(sender: str) -> str:
    """Extract the domain portion from a sender address like 'Name <user@domain.com>'."""
    match = re.search(r"@([\w.\-]+)", sender)
    return match.group(1).lower() if match else ""


def _keywords_match(text: str, keywords: list) -> bool:
    """Return True if any keyword appears in text (case-insensitive)."""
    text_lower = text.lower()
    return any(kw.lower() in text_lower for kw in keywords)


def categorize_email(em: dict, rules: dict) -> dict:
    """Apply rules to a single email and set category, priority, and note."""
    subject = em.get("subject", "")
    sender = em.get("sender", "")
    body = em.get("body", "")
    sender_domain = _extract_sender_domain(sender)

    for rule in rules.get("rules", []):
        match_cfg = rule.get("match", {})
        matched = False

        # Sender domain match
        domains = [d.lower() for d in match_cfg.get("sender_domains", [])]
        if domains and sender_domain and sender_domain in domains:
            matched = True

        # Sender address match (full or partial)
        sender_patterns = match_cfg.get("sender_contains", [])
        if not matched and sender_patterns and _keywords_match(sender, sender_patterns):
            matched = True

        # Subject keyword match
        subject_kws = match_cfg.get("subject_keywords", [])
        if not matched and subject_kws and _keywords_match(subject, subject_kws):
            matched = True

        # Body keyword match
        body_kws = match_cfg.get("body_keywords", [])
        if not matched and body_kws and _keywords_match(body, body_kws):
            matched = True

        if matched:
            em["category"] = rule.get("category", "Other")
            em["priority"] = rule.get("priority", "Medium")
            em["note"] = rule.get("note", "")
            return em

    # No rule matched — use defaults
    em["category"] = rules.get("default_category", "Other")
    em["priority"] = rules.get("default_priority", "Medium")
    em["note"] = ""
    return em


def categorize_all(emails: list, rules: dict) -> list:
    """Categorize all emails using the loaded rules."""
    print(f"Categorizing {len(emails)} emails using rules...")
    for em in emails:
        categorize_email(em, rules)
    return emails


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
ALT_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")

PRIORITY_COLORS = {
    "High": "C00000",
    "Medium": "E36C09",
    "Low": "375623",
}

COLUMNS = [
    ("Date", 20),
    ("Sender", 35),
    ("Subject", 45),
    ("Category", 22),
    ("Priority", 12),
    ("Note", 40),
    ("Body Preview", 60),
]


def _write_header(ws):
    ws.append([col for col, _ in COLUMNS])
    for col_idx, (_, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"


def _write_row(ws, row_num: int, em: dict):
    body_preview = em.get("body", "")[:200].replace("\n", " ").strip()
    values = [
        em["date"],
        em["sender"],
        em["subject"],
        em["category"],
        em["priority"],
        em.get("note", ""),
        body_preview,
    ]
    ws.append(values)
    fill = ALT_FILL if row_num % 2 == 0 else None
    priority_color = PRIORITY_COLORS.get(em.get("priority", "Medium"), "E36C09")

    for col_idx in range(1, len(values) + 1):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill:
            cell.fill = fill
        if col_idx == 5:  # Priority column
            cell.font = Font(color=priority_color, bold=True)


def export_to_excel(emails: list, output_path: str):
    """Write categorized emails to an Excel file with one sheet per category
    plus an 'All Emails' summary sheet and a Stats sheet."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    # --- All Emails sheet ---
    ws_all = wb.create_sheet("All Emails")
    _write_header(ws_all)
    for i, em in enumerate(emails, 2):
        _write_row(ws_all, i, em)

    # --- Per-category sheets ---
    categories: dict = {}
    for em in emails:
        cat = em.get("category", "Other")
        categories.setdefault(cat, []).append(em)

    for cat in sorted(categories):
        # Excel sheet names max 31 chars, strip invalid chars
        safe_name = cat[:31].translate(str.maketrans('', '', r'/\*?[]:|'))
        ws = wb.create_sheet(safe_name)
        _write_header(ws)
        for i, em in enumerate(categories[cat], 2):
            _write_row(ws, i, em)

    # --- Stats sheet ---
    ws_stats = wb.create_sheet("Stats")
    ws_stats.append(["Category", "Count", "High", "Medium", "Low"])
    for col_idx, width in [(1, 28), (2, 10), (3, 10), (4, 10), (5, 10)]:
        ws_stats.column_dimensions[get_column_letter(col_idx)].width = width
    for col_idx in range(1, 6):
        cell = ws_stats.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")

    for row_idx, cat in enumerate(sorted(categories), 2):
        items = categories[cat]
        ws_stats.append([
            cat,
            len(items),
            sum(1 for e in items if e.get("priority") == "High"),
            sum(1 for e in items if e.get("priority") == "Medium"),
            sum(1 for e in items if e.get("priority") == "Low"),
        ])
        # Bold the category name
        ws_stats.cell(row=row_idx, column=1).font = Font(bold=True)

    wb.save(output_path)
    print(f"\nExport complete: {output_path}")
    print(f"  Total emails processed : {len(emails)}")
    for cat, items in sorted(categories.items()):
        print(f"  {cat:<30} {len(items)} emails")


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def parse_args():
    parser = argparse.ArgumentParser(
        description="Filter and organize Namecheap emails into an Excel report."
    )
    parser.add_argument(
        "--limit",
        type=int,
        default=100,
        help="Maximum number of recent emails to process (default: 100).",
    )
    parser.add_argument(
        "--folder",
        default="INBOX",
        help="IMAP folder to read from (default: INBOX).",
    )
    parser.add_argument(
        "--output",
        default="email_report.xlsx",
        help="Output Excel filename (default: email_report.xlsx).",
    )
    parser.add_argument(
        "--since",
        type=int,
        default=30,
        metavar="DAYS",
        help="Only fetch emails from the last N days (default: 30).",
    )
    parser.add_argument(
        "--rules",
        default="email_rules.json",
        help="Path to rules JSON file (default: email_rules.json).",
    )
    return parser.parse_args()


def main():
    args = parse_args()

    rules = load_rules(args.rules)

    print("Connecting to mail server...")
    conn = connect_imap()
    print(f"Connected to {IMAP_HOST} as {EMAIL_ADDRESS}")

    emails = fetch_emails(conn, folder=args.folder, limit=args.limit, since_days=args.since)
    conn.logout()

    if not emails:
        print("No emails found matching criteria. Exiting.")
        return

    emails = categorize_all(emails, rules)
    export_to_excel(emails, args.output)


if __name__ == "__main__":
    main()
